"""Local parser and PostgreSQL retrieval acceptance; no model/API calls."""

import asyncio
import hashlib
import io
import json
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from pathlib import Path
from threading import Event as ThreadEvent
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import test_pm_conversation as conversation
import test_research_storage as storage
from alembic import command
from alembic.config import Config as AlembicConfig
from fastapi import UploadFile
from fastapi.testclient import TestClient
from sqlalchemy import func, select, update
from sqlalchemy import text as sql_text
from sqlalchemy.exc import IntegrityError

from dolphinlab.agents.engine import PlatformRuntime
from dolphinlab.agents.session import ResearchSession
from dolphinlab.api.schemas import CreateGoal, CreateMessage, CreateProject
from dolphinlab.application.document_context import run_source_freshness
from dolphinlab.application.documents import DocumentService, original_path, parse_isolated
from dolphinlab.application.knowledge import KnowledgeService
from dolphinlab.application.messages import MessageService
from dolphinlab.application.pm_workspace import PMWorkspace
from dolphinlab.domain.documents import SearchKnowledge
from dolphinlab.domain.research import GoalSpec
from dolphinlab.infrastructure.database import (
    ContextManifest,
    DocumentChunk,
    DocumentVersion,
    Message,
    Run,
    Session,
    now,
)
from dolphinlab.infrastructure.document_parser import parse_document
from dolphinlab.infrastructure.store import PostgresStore
from dolphinlab.main import create_app
from dolphinlab.research.staffing import freeze_profiles
from dolphinlab.research.tools import resolve_call

METHOD = "# Методика\n\n## Выручка\n\nЧистая выручка учитывает возвраты. Единица измерения RUB.\n\n## Revenue\n\nNet revenue excludes refunds.\n\n| Регион | Сумма, RUB |\n| --- | --- |\n| Север | 100 |\n| Юг | 200 |\n"


class DocumentParserTest(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory(prefix="dolphinlab-parser-")
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)

    def test_headings_tables_locators_long_text_and_repeatable_chunks(self):
        path = self.root / "method.md"
        path.write_text(
            METHOD + "\n## Длинный раздел\n\n" + "Предложение о методике. " * 500, encoding="utf-8"
        )
        first = parse_document(path, path.name)
        self.assertEqual(first, parse_document(path, path.name))
        chunks = first["chunks"]
        self.assertGreater(len(chunks), 7)
        for item in chunks:
            self.assertEqual(item["sha256"], hashlib.sha256(item["text"].encode()).hexdigest())
            self.assertIn("line_start", item["locator"])
            self.assertLessEqual(len(item["text"]), 6500)
        tables = [c for c in chunks if c["locator"]["kind"] == "table"]
        self.assertEqual(len(tables), 2)
        self.assertTrue(all("Сумма, RUB" in c["text"] for c in tables))
        self.assertNotIn("Revenue", chunks[0]["locator"]["headings"])

    def test_docx_body_heading_and_table_keep_source_indices(self):
        from docx import Document

        path = self.root / "method.docx"
        doc = Document()
        doc.add_heading("Выручка", level=1)
        doc.add_paragraph("Возвраты уменьшают чистую выручку.")
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [("Регион", "RUB"), ("Север", "100")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.save(path)
        result = parse_document(path, path.name)
        self.assertEqual(result["chunks"][0]["locator"]["headings"], ["Выручка"])
        self.assertEqual(result["chunks"][1]["locator"]["block"], 2)
        self.assertIn("RUB", result["chunks"][1]["text"])
        self.assertTrue(result["warnings"])

    def test_pdf_exact_page_and_blank_or_broken_file(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        path = self.root / "method.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 200 Td (Revenue excludes refunds) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.add_blank_page(width=300, height=300)
        writer.write(path)
        result = parse_document(path, path.name)
        self.assertEqual(result["chunks"][0]["locator"]["page"], 1)
        self.assertTrue(any("OCR" in w for w in result["warnings"]))
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(path)
        with self.assertRaisesRegex(ValueError, "OCR"):
            parse_document(path, path.name)
        path.write_bytes(b"not a PDF")
        with self.assertRaises(Exception):
            parse_document(path, path.name)

    def test_docx_zip_bomb_is_rejected_before_parsing(self):
        path = self.root / "huge.docx"
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"A" * 2_000_000)
        with self.assertRaisesRegex(ValueError, "compression ratio"):
            parse_document(path, path.name)

    def test_isolated_parser_does_not_inherit_credentials(self):
        path = self.root / "method.md"
        path.write_bytes(METHOD.encode("utf-8"))
        # Exercise the real child process, not a fake model or parser.
        result = parse_isolated(path, path.name)
        self.assertGreater(len(result["chunks"]), 0)
        with patch("dolphinlab.application.documents.subprocess.run") as call:
            call.return_value = SimpleNamespace(returncode=0, stdout=json.dumps(result))
            with patch.dict(
                "os.environ",
                {
                    "OPENAI_API_KEY": "private",
                    "DATABASE_URL": "private",
                    "EXECUTOR_TOKEN": "private",
                },
            ):
                parse_isolated(path, path.name)
            self.assertTrue(
                all(
                    k not in call.call_args.kwargs["env"]
                    for k in ["OPENAI_API_KEY", "DATABASE_URL", "EXECUTOR_TOKEN"]
                )
            )


@unittest.skipUnless(storage.TEST_URL, "Set TEST_DATABASE_URL to a disposable PostgreSQL database")
class DocumentStorageTest(unittest.TestCase):
    reset = storage.ResearchStorageTest.reset
    seed_inputs = storage.ResearchStorageTest.seed_inputs
    start = storage.ResearchStorageTest.start
    assert_status = storage.ResearchStorageTest.assert_status

    @classmethod
    def setUpClass(cls):
        storage.ResearchStorageTest.setUpClass()

    def setUp(self):
        storage.ResearchStorageTest.setUp(self)
        self.docs = DocumentService()

    def upload(
        self, content=METHOD, name="method.md", client_id=None, previous=None, project_id=None
    ):
        raw = content.encode() if isinstance(content, str) else content
        return asyncio.run(
            self.docs.upload(
                project_id or self.pid,
                UploadFile(filename=name, file=io.BytesIO(raw)),
                client_id or uuid4(),
                previous["revision"] if previous else 0,
                previous["document_id"] if previous else None,
            )
        )

    def ready(self, **kwargs):
        version = self.upload(**kwargs)
        self.assertTrue(self.docs.process_next(parse_document))
        self.assertEqual(self.docs.list(self.pid)[0]["state"], "ready")
        return version

    def search(self, query, **kwargs):
        return self.docs.search(self.pid, SearchKnowledge(query=query, kind="documents", **kwargs))[
            "results"
        ]

    def test_upload_replay_conflicts_original_and_project_scope(self):
        client_id = uuid4()
        version = self.upload(client_id=client_id)
        self.assertEqual(self.upload(client_id=client_id)["id"], version["id"])
        self.assert_status(409, lambda: self.upload("different", client_id=client_id))
        self.assertEqual(len(self.docs.list(self.pid)), 1)
        other = self.projects.create_project(CreateProject(name="Other"))["id"]
        self.assert_status(404, lambda: self.docs.download(other, version["id"]))
        self.assert_status(404, lambda: self.upload(previous=version, project_id=other))
        with Session() as db:
            record = db.get(DocumentVersion, version["id"])
            self.assertEqual(original_path(record).read_text(encoding="utf-8"), METHOD)

    def test_version_filters_ru_en_retrieval_and_exact_citations(self):
        first = self.ready()
        ru = self.search("возврат")
        en = self.search("refund")
        self.assertTrue(ru)
        self.assertTrue(en)
        source = self.docs.read(self.pid, ru[0]["source_id"], 2)
        self.assertEqual(source["sha256"], hashlib.sha256(source["text"].encode()).hexdigest())
        self.assertLessEqual(len(source["neighbors"]), 4)
        second = self.ready(
            content="# Методика\n\nВыручка включает возвраты, новое определение.", previous=first
        )
        self.assertEqual(second["revision"], 2)
        self.assertFalse(self.search("refund"))
        self.assertEqual(
            self.search("refund", include_historical=True)[0]["version_id"], first["id"]
        )
        self.assertEqual(
            self.search("refund", version_id=first["id"])[0]["version_id"], first["id"]
        )
        self.assertFalse(self.docs.read(self.pid, source["source_id"])["current"])
        self.assert_status(409, lambda: self.upload(previous=first))
        other = self.projects.create_project(CreateProject(name="Other"))["id"]
        self.assertFalse(self.docs.search(other, SearchKnowledge(query="возврат"))["results"])
        self.assert_status(404, lambda: self.docs.read(other, source["source_id"]))
        self.assertFalse(self.search("несуществующийтермин"))

    def test_failed_original_retry_and_leased_recovery(self):
        version = self.upload()
        with Session() as db:
            path = original_path(db.get(DocumentVersion, version["id"]))
        path.chmod(0o600)
        path.write_bytes(b"changed")
        self.docs.process_next(parse_document)
        self.assertEqual(self.docs.list(self.pid)[0]["state"], "failed")
        path.write_bytes(METHOD.encode("utf-8"))
        self.docs.retry(self.pid, version["id"])
        with Session.begin() as db:
            row = db.get(DocumentVersion, version["id"])
            row.state, row.lease_id, row.lease_until = (
                "processing",
                str(uuid4()),
                now() - timedelta(seconds=1),
            )
        DocumentService().process_next(parse_document)
        self.assertEqual(self.docs.list(self.pid)[0]["state"], "ready")
        self.assertFalse(self.docs.process_next(parse_document))
        with Session() as db:
            self.assertEqual(db.scalar(select(func.count()).select_from(DocumentChunk)), 4)

    def test_delete_during_processing_prevents_reappearing_chunks(self):
        version = self.upload()

        def deleted(path, name):
            result = parse_document(path, name)
            self.docs.remove(self.pid, version["document_id"], 1)
            return result

        self.docs.process_next(deleted)
        self.assertEqual(self.docs.list(self.pid)[0]["state"], "deleted")
        self.assertFalse(self.search("выручка"))
        self.assert_status(404, lambda: self.docs.retry(self.pid, version["id"]))
        self.assert_status(404, lambda: self.docs.download(self.pid, version["id"]))
        self.assertTrue(self.docs.remove(self.pid, version["document_id"], 1)["deleted"])

    def test_chunks_and_version_identity_are_database_immutable(self):
        version = self.ready()
        with self.assertRaises(IntegrityError), Session.begin() as db:
            db.execute(
                update(DocumentVersion)
                .where(DocumentVersion.id == version["id"])
                .values(sha256="0" * 64)
            )
        with self.assertRaises(IntegrityError), Session.begin() as db:
            db.execute(update(DocumentChunk).values(text="forged"))

    def test_pm_question_pins_versions_and_saves_receipts_without_run(self):
        first = self.ready()
        message = MessageService().save_message(
            self.pid, CreateMessage(content="Что такое чистая выручка?", intent="question")
        )
        workspace = PMWorkspace()
        query = {"query": "выручка", "kind": "documents"}
        result = workspace.dispatch(message["id"], "search_knowledge", query)
        self.assertTrue(result["results"])
        self.assertEqual(result, workspace.dispatch(message["id"], "search_knowledge", query))
        source_id = result["results"][0]["source_id"]
        source = workspace.dispatch(message["id"], "read_source", {"source_id": source_id})
        self.assertEqual(source["version_id"], first["id"])
        self.ready(content="# Новая версия\n\nДругое определение выручки.", previous=first)
        pinned = workspace.dispatch(
            message["id"], "search_knowledge", {"query": "возврат", "kind": "documents"}
        )
        self.assertEqual(pinned["results"][0]["version_id"], first["id"])
        self.assertEqual(self.runs.list_runs(self.pid), [])

    def test_selected_sources_freeze_run_and_survive_delete_and_rematerialization(self):
        version = self.ready()
        source = self.search("возврат")[0]
        selected = {
            "version_id": source["version_id"],
            "chunk_id": source["chunk_id"],
            "sha256": source["sha256"],
            "reason": "Definition of revenue",
        }
        goal = GoalSpec(
            question="Apply revenue definition",
            research_type="analysis",
            document_sources=[selected],
        )
        self.research.create_goal(
            self.pid, CreateGoal(client_id=uuid4(), expected_revision=1, goal=goal)
        )
        self.seed_inputs()
        run = self.start()
        with Session() as db:
            record = db.get(Run, run["id"])
            frozen = record.project_context["document_sources"][0]
            self.assertEqual(frozen["sha256"], source["sha256"])
        profile = freeze_profiles()[0]
        runtime = SimpleNamespace(
            actor=lambda a: a,
            platform_run_id=run["id"],
            staffing=SimpleNamespace(profile=lambda a: profile),
            store=PostgresStore(run["id"], Path(storage._workspace.name) / run["id"]),
        )
        evidence = asyncio.run(
            PlatformRuntime.dispatch(
                runtime, {"role": "tech_lead"}, "read_source", {"source_id": source["source_id"]}
            )
        )
        self.assertEqual(evidence["sha256"], frozen["sha256"])
        deleted = self.docs.remove(self.pid, version["document_id"], 1)
        self.assertIn(run["id"], deleted["retained_run_ids"])
        recovered = asyncio.run(
            PlatformRuntime.dispatch(
                runtime, {"role": "tech_lead"}, "read_source", {"source_id": source["source_id"]}
            )
        )
        self.assertEqual(evidence["text"], recovered["text"])
        self.assert_status(404, lambda: self.docs.read(self.pid, source["source_id"]))
        # Frozen file validation is independent of live document rows/originals.
        with tempfile.TemporaryDirectory() as directory:
            record.inputs = []
            session = ResearchSession(None, None)
            with patch("dolphinlab.agents.session.RUNS_ROOT", Path(directory)):
                task = session.materialize(record, SimpleNamespace())
                self.assertEqual(session.materialize(record, SimpleNamespace()), task)
                file = task.parent / "document-source-1.txt"
                self.assertEqual(file.read_text(encoding="utf-8"), frozen["text"])
                file.chmod(0o600)
                file.write_text("changed", encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "Frozen"):
                    session.materialize(record, SimpleNamespace())

    def test_api_upload_search_validation_and_original(self):
        with TestClient(create_app(), base_url="http://localhost") as client:
            root = f"/api/projects/{self.pid}"
            response = client.post(
                root + "/documents",
                data={"client_id": str(uuid4())},
                files={"file": ("method.md", METHOD.encode())},
            )
            self.assertEqual(response.status_code, 201, response.text)
            version = response.json()
            self.docs.process_next(parse_document)
            self.assertEqual(client.get(root + "/documents").status_code, 200)
            response = client.post(root + "/knowledge/search", json={"query": "возврат"})
            self.assertEqual(response.status_code, 200, response.text)
            source = response.json()["results"][0]
            self.assertEqual(
                client.get(root + "/knowledge/sources/" + source["source_id"]).status_code, 200
            )
            self.assertEqual(
                client.get(root + f"/documents/versions/{version['id']}/original").content,
                METHOD.encode(),
            )
            self.assertEqual(
                client.post(
                    root + "/knowledge/search", json={"query": "x", "limit": 999}
                ).status_code,
                422,
            )
            self.assertEqual(
                client.get(
                    root + "/knowledge/sources/" + source["source_id"] + "?neighbors=99"
                ).status_code,
                422,
            )

    def test_profile_v4_adds_tools_without_mutating_v3(self):
        self.assertEqual(freeze_profiles(revision=4)[0]["revision"], 4)
        old = freeze_profiles(revision=3)[0]
        self.assertNotIn("read_source", [t["function"]["name"] for t in old["tools"]])
        with self.assertRaises(ValueError):
            resolve_call(
                "tech_lead",
                "read_source",
                {"source_id": "chunk:any"},
                revision=3,
                schemas=old["tools"],
            )

    def test_document_memory_retains_quote_and_expires_on_replacement(self):
        first = self.ready()
        message = MessageService().save_message(
            self.pid, CreateMessage(content="Запомни определение из методики", intent="question")
        )
        workspace = PMWorkspace()
        source_id = self.search("возврат")[0]["source_id"]
        evidence = workspace.dispatch(message["id"], "read_source", {"source_id": source_id})
        bad = workspace.dispatch(
            message["id"],
            "remember",
            {
                "text": "forged",
                "kind": "document_statement",
                "query_id": evidence["action_id"],
                "source_quote": "missing quote",
            },
        )
        self.assertIn("error", bad)
        saved = workspace.dispatch(
            message["id"],
            "remember",
            {
                "text": "Выручка учитывает возвраты",
                "kind": "document_statement",
                "scope": "project",
                "query_id": evidence["action_id"],
                "source_quote": "Чистая выручка учитывает возвраты.",
            },
        )
        self.assertNotIn("error", saved)
        memory = KnowledgeService().list_memory(self.pid)
        self.assertEqual(memory[0]["document_refs"][0]["version_id"], first["id"])
        self.ready(
            content="# Методика v2\n\nВозвраты исключены из определения выручки.", previous=first
        )
        self.assertFalse(KnowledgeService().list_memory(self.pid))
        history = KnowledgeService().list_memory(self.pid, include_stale=True)
        self.assertFalse(history[0]["current"])
        self.assertEqual(history[0]["status"], "reported")
        self.assertFalse(
            self.docs.search(self.pid, SearchKnowledge(query="возврат", kind="memory"))["results"]
        )
        self.assertTrue(
            self.docs.search(
                self.pid, SearchKnowledge(query="возврат", kind="memory", include_historical=True)
            )["results"]
        )

    def test_scripted_pm_cites_document_without_research(self):
        self.ready()
        conversation.ConversationTest.configure(self)
        source_id = self.search("возврат")[0]["source_id"]
        saved = MessageService().save_message(
            self.pid, CreateMessage(content="Как учитываются возвраты?", intent="question")
        )
        with Session() as db:
            message = db.get(Message, saved["id"])
        responses = [
            conversation.tool(
                "search_knowledge", {"query": "возврат", "kind": "documents"}, "search-doc"
            ),
            conversation.tool("read_source", {"source_id": source_id, "neighbors": 0}, "read-doc"),
            {"role": "assistant", "content": "По методике возвраты учтены. Источник: " + source_id},
        ]
        factory = conversation.ConversationTest.handle(self, message, responses)
        self.assertEqual(len(factory.client.calls), 3)
        self.assertEqual(self.runs.list_runs(self.pid), [])
        messages = MessageService().list_messages(self.pid)
        self.assertIn(source_id, messages[-1]["content"])

    def test_ingestion_commit_failure_rolls_back_all_chunks_and_retries_once(self):
        version = self.upload()
        # Crash after some INSERTs have been flushed, before committing the version.
        original = Session.class_.flush

        def interrupted(session, objects=None):
            inserting = any(isinstance(x, DocumentChunk) for x in session.new)
            result = original(session, objects)
            if inserting:
                raise RuntimeError("interrupted commit")
            return result

        with patch.object(Session.class_, "flush", interrupted):
            with self.assertRaisesRegex(RuntimeError, "interrupted"):
                self.docs.process_next(parse_document)
        with Session.begin() as db:
            self.assertEqual(db.scalar(select(func.count()).select_from(DocumentChunk)), 0)
            row = db.get(DocumentVersion, version["id"])
            self.assertEqual(row.state, "processing")
            row.lease_until = now() - timedelta(seconds=1)
        DocumentService().process_next(parse_document)
        self.assertEqual(self.docs.list(self.pid)[0]["state"], "ready")
        with Session() as db:
            ids = list(db.scalars(select(DocumentChunk.id)))
        self.assertFalse(self.docs.process_next(parse_document))
        with Session() as db:
            self.assertEqual(ids, list(db.scalars(select(DocumentChunk.id))))

    def test_invalid_parser_output_and_repeated_crashes_are_visible_failures(self):
        version = self.upload()
        self.docs.process_next(
            lambda *args: {
                "chunks": [{"text": "forged", "sha256": "0" * 64, "locator": {}}],
                "warnings": [],
            }
        )
        self.assertEqual(self.docs.list(self.pid)[0]["state"], "failed")
        with Session.begin() as db:
            row = db.get(DocumentVersion, version["id"])
            row.state, row.attempts, row.lease_until = "processing", 3, now() - timedelta(seconds=1)
        self.docs.process_next(
            lambda *args: self.fail("Repeated crashes must not reparse automatically")
        )
        self.assertEqual(self.docs.list(self.pid)[0]["state"], "failed")

    def test_retrieval_baseline_expected_passages_in_top_five(self):
        self.ready()
        cases = [
            ("выручка возвраты", "Чистая выручка"),
            ("refund", "Net revenue"),
            ("единица RUB", "Единица измерения"),
            ("Север", "100"),
            ("Юг", "200"),
        ]
        for query, expected in cases:
            with self.subTest(query=query):
                self.assertTrue(any(expected in h["text"] for h in self.search(query, limit=5)))
        self.assertFalse(self.search("каузальная скидка"))

    def test_project_deletion_cleans_document_rows_and_readonly_originals(self):
        version = self.ready()
        with Session() as db:
            path = original_path(db.get(DocumentVersion, version["id"]))
        self.projects.delete_project(self.pid)
        self.assertFalse(path.exists())
        with Session() as db:
            self.assertEqual(db.scalar(select(func.count()).select_from(DocumentVersion)), 0)
            self.assertEqual(db.scalar(select(func.count()).select_from(DocumentChunk)), 0)

    def test_project_deletion_during_extraction_does_not_crash_worker(self):
        self.upload()

        def removed(path, name):
            result = parse_document(path, name)
            self.projects.delete_project(self.pid)
            return result

        self.assertTrue(self.docs.process_next(removed))
        with Session() as db:
            self.assertEqual(db.scalar(select(func.count()).select_from(DocumentVersion)), 0)

    def test_concurrent_ingestion_claims_only_one_lease(self):
        self.upload()
        entered, release = ThreadEvent(), ThreadEvent()

        def slow(path, name):
            entered.set()
            if not release.wait(10):
                raise RuntimeError("Test synchronization timed out")
            return parse_document(path, name)

        with ThreadPoolExecutor(max_workers=2) as pool:
            future = pool.submit(self.docs.process_next, slow)
            self.assertTrue(entered.wait(5))
            try:
                self.assertFalse(DocumentService().process_next(parse_document))
            finally:
                release.set()
            self.assertTrue(future.result())
        version = self.docs.list(self.pid)[0]
        self.assertEqual(version["state"], "ready")
        self.assertEqual(version["attempts"], 1)

    def test_upgrade_0006_backfills_memory_without_rewriting_its_text(self):
        self.reset("0006")
        with storage.engine.begin() as db:
            db.execute(
                sql_text(
                    "INSERT INTO projects(id,name,description,status,archived,created) VALUES ('legacy','Legacy','','draft',false,now())"
                )
            )
            db.execute(
                sql_text(
                    "INSERT INTO messages(id,project_id,role,content,directive,status,created,intent,context_binding) VALUES ('m','legacy','user','text',false,'delivered',now(),'question','{}')"
                )
            )
            db.execute(
                sql_text(
                    "INSERT INTO project_memory(id,project_id,message_id,scope,kind,text,status,input_ids,created) VALUES ('memory','legacy','m','project','hypothesis','ЕДИНИЦА RUB','reported','[]',now())"
                )
            )
        config = AlembicConfig(str(Path(__file__).parents[1] / "alembic.ini"))
        command.upgrade(config, "head")
        with storage.engine.connect() as db:
            row = db.execute(
                sql_text(
                    "SELECT text,search_text,document_refs FROM project_memory WHERE id='memory'"
                )
            ).one()
            self.assertEqual(tuple(row), ("ЕДИНИЦА RUB", "единица rub", []))
        command.downgrade(config, "0006")
        with storage.engine.connect() as db:
            self.assertEqual(
                db.scalar(sql_text("SELECT text FROM project_memory WHERE id='memory'")),
                "ЕДИНИЦА RUB",
            )

    def test_model_context_is_bounded_but_manifest_keeps_all_pinned_versions(self):
        for index in range(21):
            self.ready(
                content=f"# Методика {index}\n\nОпределение выручки {index}.",
                name=f"method-{index}.md",
            )
        message = MessageService().save_message(
            self.pid, CreateMessage(content="Найди определение", intent="question")
        )
        context = KnowledgeService().context(message["id"])
        self.assertEqual(len(context["documents"]), 20)
        self.assertEqual(context["document_scope"]["pinned_versions"], 21)
        self.assertNotIn("documents", context["source_message"]["binding"])
        with Session() as db:
            manifest = db.get(ContextManifest, context["context_manifest_id"])
            self.assertEqual(len(manifest.snapshot["source_message"]["binding"]["documents"]), 21)

    def test_message_pins_queued_upload_then_can_read_same_version_when_ready(self):
        version = self.upload()
        saved = MessageService().save_message(
            self.pid, CreateMessage(content="Как методика определяет выручку?", intent="question")
        )
        context = KnowledgeService().context(saved["id"])
        self.assertEqual(context["documents"][0]["state"], "queued")
        self.assertEqual(context["documents"][0]["version_id"], version["id"])
        self.docs.process_next(parse_document)
        result = PMWorkspace().dispatch(
            saved["id"], "search_knowledge", {"query": "возврат", "kind": "documents"}
        )
        self.assertEqual(result["results"][0]["version_id"], version["id"])

    def test_document_replacement_changes_applicability_without_rewriting_run(self):
        version = self.ready()
        source = self.search("возврат")[0]
        goal = GoalSpec(
            question="Use method v1",
            document_sources=[
                {
                    "version_id": source["version_id"],
                    "chunk_id": source["chunk_id"],
                    "sha256": source["sha256"],
                    "reason": "Definition",
                }
            ],
        )
        self.research.create_goal(
            self.pid, CreateGoal(client_id=uuid4(), expected_revision=1, goal=goal)
        )
        self.seed_inputs()
        started = self.start()
        with Session() as db:
            run = db.get(Run, started["id"])
            frozen = run.goal_snapshot
            self.assertTrue(run_source_freshness(db, run)["current"])
        self.ready(
            content="# Новая методика\n\nВыручка без НДС неизвестна без vat_amount.",
            previous=version,
        )
        with Session() as db:
            run = db.get(Run, started["id"])
            self.assertEqual(
                run_source_freshness(db, run)["stale_document_version_ids"], [version["id"]]
            )
            self.assertEqual(run.goal_snapshot, frozen)
